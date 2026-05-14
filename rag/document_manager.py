"""
Document manager: handles PDF, DOCX, TXT, MD uploads.
Chunks documents and stores them for retrieval.
"""
import os
import hashlib
import logging
from typing import List, Dict, Optional, Tuple
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class DocumentChunk:
    chunk_id: str
    doc_id: str
    text: str
    metadata: Dict = field(default_factory=dict)


@dataclass
class Document:
    doc_id: str
    filename: str
    user_id: str
    file_type: str
    chunks: List[DocumentChunk] = field(default_factory=list)
    metadata: Dict = field(default_factory=dict)


class DocumentManager:
    """
    Manages user-uploaded documents.
    Supports PDF, DOCX, TXT, MD, HTML.
    Chunks text and deduplicates via content hash.
    """

    SUPPORTED_TYPES = {".pdf", ".docx", ".txt", ".md", ".html"}

    def __init__(
        self,
        chunk_size: int = 512,
        chunk_overlap: int = 64,
        storage_dir: str = "data/documents",
    ):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        # doc_id -> Document
        self._store: Dict[str, Document] = {}
        # content hash -> doc_id (dedup)
        self._hash_index: Dict[str, str] = {}

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def add_document(
        self,
        file_path: str,
        user_id: str,
        extra_metadata: Optional[Dict] = None,
    ) -> Document:
        """Parse a file, chunk it, and store it. Returns the Document."""
        path = Path(file_path)
        suffix = path.suffix.lower()
        if suffix not in self.SUPPORTED_TYPES:
            raise ValueError(
                f"Unsupported file type '{suffix}'. "
                f"Supported: {self.SUPPORTED_TYPES}"
            )

        raw_text = self._extract_text(path)
        content_hash = hashlib.sha256(raw_text.encode()).hexdigest()

        # Deduplication
        if content_hash in self._hash_index:
            existing_id = self._hash_index[content_hash]
            logger.info(
                "Document '%s' is a duplicate of '%s'. Skipping.",
                path.name,
                existing_id,
            )
            return self._store[existing_id]

        doc_id = f"{user_id}_{content_hash[:16]}"
        chunks = self._chunk_text(raw_text, doc_id, path.name)

        doc = Document(
            doc_id=doc_id,
            filename=path.name,
            user_id=user_id,
            file_type=suffix,
            chunks=chunks,
            metadata={
                "content_hash": content_hash,
                "char_count": len(raw_text),
                **(extra_metadata or {}),
            },
        )

        self._store[doc_id] = doc
        self._hash_index[content_hash] = doc_id
        logger.info("Added document '%s' (%d chunks).", path.name, len(chunks))
        return doc

    def get_document(self, doc_id: str) -> Optional[Document]:
        return self._store.get(doc_id)

    def list_documents(self, user_id: Optional[str] = None) -> List[Document]:
        docs = list(self._store.values())
        if user_id:
            docs = [d for d in docs if d.user_id == user_id]
        return docs

    def delete_document(self, doc_id: str) -> bool:
        doc = self._store.pop(doc_id, None)
        if doc is None:
            return False
        h = doc.metadata.get("content_hash")
        if h:
            self._hash_index.pop(h, None)
        logger.info("Deleted document '%s'.", doc_id)
        return True

    def get_all_chunks(self, user_id: Optional[str] = None) -> List[DocumentChunk]:
        docs = self.list_documents(user_id)
        chunks: List[DocumentChunk] = []
        for doc in docs:
            chunks.extend(doc.chunks)
        return chunks

    # ------------------------------------------------------------------
    # Internals
    # ------------------------------------------------------------------

    def _extract_text(self, path: Path) -> str:
        suffix = path.suffix.lower()
        try:
            if suffix == ".pdf":
                return self._extract_pdf(path)
            elif suffix == ".docx":
                return self._extract_docx(path)
            elif suffix in {".txt", ".md"}:
                return path.read_text(encoding="utf-8", errors="replace")
            elif suffix == ".html":
                return self._extract_html(path)
        except Exception as exc:
            logger.warning("Text extraction failed for '%s': %s", path, exc)
            return ""
        return ""

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            import pdfminer.high_level as pdf_hl  # type: ignore
            return pdf_hl.extract_text(str(path))
        except ImportError:
            pass
        try:
            import PyPDF2  # type: ignore
            text_parts = []
            with open(path, "rb") as fh:
                reader = PyPDF2.PdfReader(fh)
                for page in reader.pages:
                    text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except ImportError:
            logger.warning("No PDF library found (pdfminer or PyPDF2). Returning empty.")
            return ""

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            import docx  # type: ignore
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            logger.warning("python-docx not installed. Returning empty for docx.")
            return ""

    @staticmethod
    def _extract_html(path: Path) -> str:
        raw = path.read_text(encoding="utf-8", errors="replace")
        try:
            from html.parser import HTMLParser

            class _Strip(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self._parts: List[str] = []

                def handle_data(self, data: str):
                    self._parts.append(data)

            parser = _Strip()
            parser.feed(raw)
            return " ".join(parser._parts)
        except Exception:
            return raw

    def _chunk_text(
        self, text: str, doc_id: str, filename: str
    ) -> List[DocumentChunk]:
        """Sliding-window character chunking with overlap."""
        chunks: List[DocumentChunk] = []
        start = 0
        idx = 0
        while start < len(text):
            end = start + self.chunk_size
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunk_id = f"{doc_id}_chunk_{idx}"
                chunks.append(
                    DocumentChunk(
                        chunk_id=chunk_id,
                        doc_id=doc_id,
                        text=chunk_text,
                        metadata={"source_file": filename, "chunk_index": idx},
                    )
                )
                idx += 1
            start += self.chunk_size - self.chunk_overlap
        return chunks
